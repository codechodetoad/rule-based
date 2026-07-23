#!/usr/bin/env python3
"""Bệnh (ICD-10) — parse file Excel gốc của Bộ Y tế -> library/diseases.json

Nguồn: 06-byt-kem.xlsx (Phụ lục Danh mục mã bệnh theo ICD-10, TT 06/2026/TT-BYT).
File gồm 4 sheet nối tiếp nhau (Table 1..4); chỉ Table 1 có 3 dòng tiêu đề/số cột ở đầu.
Cột dùng: MÃ BỆNH (18), DISEASE NAME WHO EN (20), TÊN BỆNH VN (22)  [1-indexed].

Output: dict {tên (lowercase) -> mã}, KHÔNG dùng list/object/array lồng nhau ->
nhẹ nhất có thể, tên là key nên tra cứu O(1), mã là string thuần (nối "," nếu
1 tên trùng nhau ứng với nhiều mã khác nhau — hiếm, không mất thông tin):
  {"bệnh tả": "A00", "cholera": "A00", ...}

Chạy: python library/build/parse_icd_xlsx.py
"""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
XLSX = "C:/Users/Lenovo/Documents/iLovePDF_Output/06-byt-kem.xlsx"
OUT = os.path.join(ROOT, "library", "diseases.json")

COL_CODE = 17   # 0-indexed -> "MÃ BỆNH"
COL_EN = 19     # "DISEASE NAME WHO 2019 (ENGLISH)"
COL_VN = 21     # "TÊN BỆNH"


def clean_text(s):
    if not s:
        return ""
    s = str(s).replace("\n", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def clean_code(c):
    # Mã có thể mang hậu tố '*' (mã biểu hiện) hoặc '†' (mã nguyên nhân) của hệ
    # thống lưỡng phân dagger-asterisk -> giữ lại code sạch để so khớp, ký hiệu
    # phân loại không phải một phần của mã bệnh.
    c = str(c).strip()
    c = c.replace("*", "").replace("†", "").strip()
    return c


def iter_rows(wb):
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = ws.iter_rows(values_only=True)
        if sn == "Table 1":
            next(rows, None)
            next(rows, None)
            next(rows, None)
        for row in rows:
            yield row


def add_entry(d, text, code):
    key = text.lower()
    if key in d:
        codes = d[key].split(",")
        if code not in codes:
            d[key] = d[key] + "," + code
    else:
        d[key] = code


def main():
    if not os.path.exists(XLSX):
        print(f"KHÔNG TÌM THẤY: {XLSX}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    out = {}
    n_rows = 0
    n_codes = set()

    for row in iter_rows(wb):
        if row is None or len(row) <= COL_VN:
            continue
        raw_code = row[COL_CODE]
        if raw_code is None:
            continue
        code = clean_code(raw_code)
        if not code:
            continue
        n_rows += 1
        n_codes.add(code)

        name_vn = clean_text(row[COL_VN])
        name_en = clean_text(row[COL_EN])

        for name in (name_vn, name_en):
            if name:
                add_entry(out, name, code)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"Đọc {n_rows} dòng, {len(n_codes)} mã ICD-10 unique")
    print(f"Ghi {len(out)} key -> {OUT}")


if __name__ == "__main__":
    main()
